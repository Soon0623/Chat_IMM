import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from pptx import Presentation
from bridge import build
from api import search

class Integration(unittest.TestCase):
    def test_lifecycle(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d); repo=root/'repo'; src=repo/'submissions/student/demo'; src.mkdir(parents=True)
            prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text='자성 로봇'
            slide.notes_slide.notes_text_frame.text='독립 자화 방향'
            table=slide.shapes.add_table(1,1,0,0,1000000,1000000).table
            table.cell(0,0).text='표 데이터'
            prs.save(src/'demo.pptx')
            desc=src/'description.md'; desc.write_text('특수 100% 자성 설명',encoding='utf-8')
            vault=root/'vault'; db=root/'search.db'
            build(repo,vault,db)
            with patch.dict(os.environ,{'BRIDGE_DB':str(db),'BRIDGE_API_KEY':'key'}):
                self.assertTrue(search('독립',5,'key')['results'])
                self.assertTrue(search('표 데이터',5,'key')['results'])
                self.assertEqual(len(search('100%',5,'key')['results']),1)
            personal=vault/'personal.md'; personal.write_text('개인 메모')
            desc.write_text('변경된 설명',encoding='utf-8'); build(repo,vault,db)
            with sqlite3.connect(db) as c:
                self.assertTrue(c.execute("select count(*) from chunks where text like '%변경된%'").fetchone()[0])
            before=db.read_bytes(); (src/'broken.pptx').write_bytes(b'bad')
            with self.assertRaises(Exception): build(repo,vault,db)
            self.assertEqual(before,db.read_bytes()); (src/'broken.pptx').unlink()
            (src/'demo.pptx').unlink(); desc.unlink(); build(repo,vault,db)
            self.assertEqual(list((vault/'GitHub_Auto').glob('*.md')),[])
            self.assertTrue(personal.exists())
            with sqlite3.connect(db) as c: self.assertEqual(c.execute('select count(*) from chunks').fetchone()[0],0)

if __name__=='__main__': unittest.main()
