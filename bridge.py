"""Deterministic PPTX -> Markdown + SQLite retrieval snapshot."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import sqlite3
import tempfile
from urllib.parse import quote
from pptx import Presentation


def shapes_text(shapes):
    parts = []
    for shape in shapes:
        if shape.shape_type == 6:
            parts.extend(shapes_text(shape.shapes))
        elif shape.has_table:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
        elif shape.has_text_frame:
            parts.append(shape.text_frame.text)
    return [s.strip() for s in parts if s.strip()]


def extract(path):
    if path.suffix.lower() == ".pptx":
        prs = Presentation(str(path))
        sections = []
        for i, slide in enumerate(prs.slides, 1):
            text = "\n\n".join(shapes_text(slide.shapes))
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame and frame.text.strip():
                    text += "\n\n발표자 노트:\n" + frame.text
            sections.append((i, text or "[추출할 텍스트 없음: 이미지 자료 확인 필요]"))
        return sections
    return [(None, path.read_text(encoding="utf-8-sig"))]


def build(repo, vault, db, repo_url="", branch="main"):
    repo, vault, db = Path(repo).resolve(), Path(vault).resolve(), Path(db).resolve()
    source = repo / "submissions"
    if not source.is_dir():
        raise ValueError("submissions 폴더가 없습니다. 기존 출력은 보존합니다.")
    vault.mkdir(parents=True, exist_ok=True)
    db.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix="bridge-", dir=vault))
    dbtmp = None
    try:
        records = []
        for path in sorted(source.rglob("*")):
            if not path.is_file():
                continue
            if path.is_symlink():
                raise ValueError(f"심볼릭 링크는 지원하지 않습니다: {path}")
            if path.suffix.lower() == ".ppt":
                raise ValueError(f"PPT를 PPTX로 변환해 주세요: {path.name}")
            if path.suffix.lower() not in {".pptx", ".md", ".txt"}:
                continue
            rel = path.relative_to(repo).as_posix()
            key = hashlib.sha256(rel.encode()).hexdigest()[:20]
            url = (repo_url.rstrip("/") + "/blob/" + quote(branch, safe="") + "/" + quote(rel)) if repo_url else rel
            sections = extract(path)
            body = "\n\n".join((f"## 슬라이드 {n}\n\n" if n else "") + t for n, t in sections)
            note = f"---\nsource: {json.dumps(rel, ensure_ascii=False)}\n---\n\n# {path.stem}\n\n출처: {url}\n\n{body}\n"
            (stage / f"{key}.md").write_text(note, encoding="utf-8")
            for slide, text in sections:
                # Overlapping chunks retain source and slide provenance.
                for offset in range(0, len(text), 1000):
                    chunk = text[offset:offset + 1200]
                    records.append((f"{key}:{slide}:{offset}", rel, slide, url, chunk))
        fd, name = tempfile.mkstemp(prefix="index-", suffix=".db", dir=db.parent)
        os.close(fd)
        dbtmp = Path(name)
        with sqlite3.connect(dbtmp) as conn:
            conn.execute("CREATE TABLE chunks (id TEXT PRIMARY KEY, source TEXT, slide INTEGER, url TEXT, text TEXT)")
            conn.executemany("INSERT INTO chunks VALUES (?,?,?,?,?)", records)
        # Only owned generated files are removed; personal vault notes are untouched.
        target = vault / "GitHub_Auto"
        target.mkdir(exist_ok=True)
        for note in stage.glob("*.md"):
            os.replace(note, target / note.name)
        expected = {f"{row[0].split(':')[0]}.md" for row in records}
        for old in target.glob("*.md"):
            if len(old.stem) == 20 and all(c in "0123456789abcdef" for c in old.stem) and old.name not in expected:
                old.unlink()
        os.replace(dbtmp, db)
        print(f"Indexed {len(records)} chunks; notes: {target}")
    finally:
        shutil.rmtree(stage, ignore_errors=True)
        if dbtmp:
            dbtmp.unlink(missing_ok=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", default=".")
    parser.add_argument("--vault", default="vault")
    parser.add_argument("--db", default="data/search.db")
    parser.add_argument("--repo-url", default="")
    parser.add_argument("--branch", default="main")
    args = parser.parse_args()
    build(args.repo, args.vault, args.db, args.repo_url, args.branch)
